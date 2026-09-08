#!/usr/bin/env python3
"""Temporary B2 maintenance script: acquire a price-free normalized TSB snapshot.

This file is intentionally deleted before Phase B2 merges. It reads the official
monthly TSB workbook in memory and writes only Bizzat's normalized source facts:
vehicle code, brand/type text, and available model years. Kasko values are never
written to the output.
"""

from __future__ import annotations

import argparse
import io
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

BASE_URL = "https://www.tsb.org.tr"
MONTHS_URL = f"{BASE_URL}/InsuranceData/GetMonthList"
ARCHIVE_URL = f"{BASE_URL}/InsuranceData/GetInsuranceDataArchiveFile"
SOURCE_PAGE = f"{BASE_URL}/tr/kasko-arsiv-listesi"
HEADER_SCAN_ROWS = 25


def fold(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("ı", "i").replace("İ", "I")
    decomposed = unicodedata.normalize("NFKD", text)
    return " ".join(
        "".join(ch for ch in decomposed if not unicodedata.combining(ch))
        .lower()
        .split()
    )


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(unicodedata.normalize("NFKC", str(value)).strip().split())


def identity(value: str) -> str:
    return fold(value).upper()


def as_int(value: Any) -> int | None:
    text = clean_text(value)
    if not text:
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        return None


def unwrap(payload: Any) -> Any:
    if isinstance(payload, dict) and "Result" in payload:
        if payload.get("HasError") is True:
            raise RuntimeError(f"TSB returned an error: {payload.get('Message')}")
        return payload["Result"]
    return payload


def request_json(session: requests.Session, url: str, **params: Any) -> Any:
    response = session.get(url, params=params or None, timeout=30)
    response.raise_for_status()
    try:
        return unwrap(response.json())
    except ValueError as error:
        raise RuntimeError(f"Expected JSON from {response.url}") from error


def resolve_archive_url(session: requests.Session, year: int, month: int) -> str:
    months = request_json(session, MONTHS_URL)
    if not isinstance(months, list):
        raise RuntimeError("TSB month list response is not an array")

    month_id: int | None = None
    for item in months:
        if not isinstance(item, dict):
            continue
        try:
            order = int(item.get("MonthOrder"))
        except (TypeError, ValueError):
            continue
        if order == month:
            try:
                month_id = int(item["Id"])
            except (KeyError, TypeError, ValueError) as error:
                raise RuntimeError("TSB month item has no usable Id") from error
            break

    if month_id is None:
        raise RuntimeError(f"TSB does not expose calendar month {month}")

    payload = request_json(session, ARCHIVE_URL, Year=year, MonthId=month_id)
    if not isinstance(payload, str) or not payload.strip():
        raise RuntimeError(f"TSB archive path missing for {year}-{month:02d}")
    path = payload.strip()
    return path if path.startswith("http") else f"{BASE_URL}/{path.lstrip('/')}"


FIELD_KEYWORDS: dict[str, tuple[str, ...]] = {
    "brand_code": ("marka kodu", "markakodu", "brand code"),
    "model_code": ("model kodu", "modelkodu", "tip kodu", "tipkodu", "model code"),
    "model_year": ("model yili", "modelyili", "model year", "yil"),
    "brand": ("marka", "brand"),
    "model": ("model", "tip", "arac", "vehicle"),
    "amount": ("kasko bedeli", "kaskobedeli", "bedel", "deger", "tutar", "amount", "value"),
}
FIELD_PRIORITY = ("brand_code", "model_code", "model_year", "amount", "brand", "model")
COMBINED = "brand_model"


def match_header(value: Any) -> str | None:
    header = fold(value)
    if not header:
        return None
    if "marka" in header and ("model" in header or "tip" in header):
        return COMBINED
    for field in FIELD_PRIORITY:
        if any(keyword in header for keyword in FIELD_KEYWORDS[field]):
            return field
    return None


def locate_header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    best_index = -1
    best: dict[str, int] = {}
    for row_index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        mapping: dict[str, int] = {}
        for column_index, value in enumerate(row):
            field = match_header(value)
            if field and field not in mapping:
                mapping[field] = column_index
        if len(mapping) > len(best):
            best_index, best = row_index, mapping

    required = {"brand_code", "model_code", "model_year"}
    if best_index < 0 or not required.issubset(best):
        raise RuntimeError(f"Could not locate required TSB workbook columns; found {sorted(best)}")
    if COMBINED not in best and not {"brand", "model"}.issubset(best):
        raise RuntimeError("TSB workbook has no usable brand/model text columns")
    return best_index, best


def split_combined(value: Any) -> tuple[str, str]:
    text = clean_text(value)
    if not text:
        return "", ""
    for separator in (" - ", "-", "/"):
        if separator in text:
            head, tail = text.split(separator, 1)
            if head.strip() and tail.strip():
                return clean_text(head), clean_text(tail)
    return "", text


def cell(row: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def parse_workbook(content: bytes, *, year: int, month: int) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises several workbook-specific exceptions
        raise RuntimeError(f"Could not open TSB workbook: {error}") from error

    try:
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    header_index, columns = locate_header(rows)
    grouped: dict[str, dict[str, Any]] = {}
    years_by_key: dict[str, set[int]] = defaultdict(set)

    for row in rows[header_index + 1 :]:
        brand_code = as_int(cell(row, columns.get("brand_code")))
        model_code = as_int(cell(row, columns.get("model_code")))
        model_year = as_int(cell(row, columns.get("model_year")))
        if brand_code is None or model_code is None or model_year is None:
            continue
        if model_year < 1886 or model_year > year + 1:
            continue

        if COMBINED in columns:
            brand_raw, type_raw = split_combined(cell(row, columns.get(COMBINED)))
        else:
            brand_raw = clean_text(cell(row, columns.get("brand")))
            type_raw = clean_text(cell(row, columns.get("model")))
        if not brand_raw or not type_raw:
            continue

        source_key = f"{brand_code}-{model_code}"
        existing = grouped.get(source_key)
        if existing is None:
            grouped[source_key] = {
                "sourceKey": source_key,
                "brandRaw": brand_raw,
                "typeRaw": type_raw,
            }
        elif (
            identity(existing["brandRaw"]) != identity(brand_raw)
            or identity(existing["typeRaw"]) != identity(type_raw)
        ):
            raise RuntimeError(
                f"Conflicting brand/type identity for TSB code {source_key}: "
                f"{existing['brandRaw']} / {existing['typeRaw']} vs {brand_raw} / {type_raw}"
            )
        years_by_key[source_key].add(model_year)

    if not grouped:
        raise RuntimeError("TSB workbook produced zero normalized vehicle records")

    records = []
    for source_key in sorted(grouped):
        row = grouped[source_key]
        records.append(
            {
                **row,
                "availableModelYears": sorted(years_by_key[source_key]),
            }
        )
    return records


def acquire(year: int, month: int) -> dict[str, Any]:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": "Bizzat vehicle-catalog maintenance/1.0",
            "Accept-Language": "tr-TR,tr;q=0.9,en;q=0.8",
            "Referer": SOURCE_PAGE,
        }
    )
    archive_url = resolve_archive_url(session, year, month)
    response = session.get(archive_url, timeout=60)
    response.raise_for_status()
    records = parse_workbook(response.content, year=year, month=month)
    return {
        "provider": {
            "code": "tsb-kasko",
            "sourceName": "Türkiye Sigorta Birliği Kasko Değer Listesi",
            "sourceUrl": SOURCE_PAGE,
            "version": f"{year}-{month:02d}",
        },
        "records": records,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--year", type=int, required=True)
    parser.add_argument("--month", type=int, choices=range(1, 13), required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    snapshot = acquire(args.year, args.month)
    serialized = json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    forbidden = ("amount", "price", "bedel", "deger", "tutar")
    lowered = serialized.lower()
    if any(f'"{term}"' in lowered for term in forbidden):
        raise RuntimeError("Normalized snapshot unexpectedly contains a price/value field")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(serialized + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "version": snapshot["provider"]["version"],
                "records": len(snapshot["records"]),
                "output": str(args.output),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"acquisition failed: {error}", file=sys.stderr)
        raise
